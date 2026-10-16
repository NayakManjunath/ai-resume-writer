from pathlib import Path

from docx import Document

from ai_resume_writer.contracts.models import (
    CandidateProfile,
    ResumeFormatChoice,
    ResumeImprovement,
    TargetRole,
)
from ai_resume_writer.documents.generator import (
    DocumentGenerationError,
    generate_resume_documents,
)


def _build_improvement(
    experience_years: float | None = 3.0,
    format_choice: ResumeFormatChoice = ResumeFormatChoice.USE_TEMPLATE,
) -> tuple[ResumeImprovement, CandidateProfile]:
    profile = CandidateProfile(
        name="Test Candidate",
        professional_summary="Data Scientist with Python experience.",
        skills=["Python", "SQL"],
        experience=["Data Scientist - Example Company"],
        projects=["Recommendation System"],
        education=["B.Tech"],
        certifications=["AWS Certification"],
        experience_years=experience_years,
    )

    improvement = ResumeImprovement(
        target_role=TargetRole(
            designation="Data Scientist",
            source="job_description",
        ),
        format_choice=format_choice,
        original_profile=profile,
        improved_sections={
            "Professional Summary": [
                "Data Scientist with Python experience."
            ],
            "Skills": [
                "Python",
                "SQL",
            ],
            "Experience": [
                "Data Scientist - Example Company",
            ],
            "Projects": [
                "Recommendation System",
            ],
            "Education": [
                "B.Tech",
            ],
            "Certifications": [
                "AWS Certification",
            ],
        },
        improvement_summary=[
            "Improved resume alignment."
        ],
    )

    return improvement, profile


def test_generate_resume_documents(tmp_path):
    improvement, profile = _build_improvement()

    result = generate_resume_documents(
        improvement,
        profile,
        tmp_path,
    )

    assert result.docx_path.exists()
    assert result.pdf_path.exists()

    assert result.page_validation.page_count >= 1
    assert result.page_validation.expected_max_pages == 1


def test_generated_docx_contains_resume_content(tmp_path):
    improvement, profile = _build_improvement()

    result = generate_resume_documents(
        improvement,
        profile,
        tmp_path,
    )

    document = Document(result.docx_path)

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Test Candidate" in text
    assert "PROFESSIONAL SUMMARY" in text
    assert "Python" in text
    assert "Recommendation System" in text


def test_one_page_requirement_is_validated(tmp_path):
    improvement, profile = _build_improvement(
        experience_years=2.0,
    )

    result = generate_resume_documents(
        improvement,
        profile,
        tmp_path,
    )

    assert result.page_validation.expected_max_pages == 1
    assert result.page_validation.is_valid is True


def test_candidates_above_four_years_have_no_strict_one_page_limit(
    tmp_path,
):
    improvement, profile = _build_improvement(
        experience_years=5.0,
    )

    result = generate_resume_documents(
        improvement,
        profile,
        tmp_path,
    )

    assert result.page_validation.expected_max_pages is None
    assert result.page_validation.is_valid is True


def test_missing_experience_years_has_no_strict_page_limit(tmp_path):
    improvement, profile = _build_improvement(
        experience_years=None,
    )

    result = generate_resume_documents(
        improvement,
        profile,
        tmp_path,
    )

    assert result.page_validation.expected_max_pages is None


def test_empty_sections_are_not_written(tmp_path):
    improvement, profile = _build_improvement()

    improvement = improvement.model_copy(
        update={
            "improved_sections": {
                "Skills": ["Python"],
                "Projects": [],
            }
        }
    )

    result = generate_resume_documents(
        improvement,
        profile,
        tmp_path,
    )

    document = Document(result.docx_path)

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "SKILLS" in text
    assert "PROJECTS" not in text


def test_existing_docx_can_be_used_as_baseline(tmp_path):
    source_path = tmp_path / "existing_resume.docx"

    source_document = Document()
    source_document.add_paragraph("Original Resume")
    source_document.save(source_path)

    improvement, profile = _build_improvement(
        format_choice=ResumeFormatChoice.KEEP_EXISTING,
    )

    result = generate_resume_documents(
        improvement,
        profile,
        tmp_path / "generated",
        existing_docx_path=source_path,
    )

    assert result.docx_path.exists()
    assert result.pdf_path.exists()

    document = Document(result.docx_path)

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Original Resume" not in text
    assert "Test Candidate" in text


def test_generation_creates_output_directory(tmp_path):
    improvement, profile = _build_improvement()

    output_dir = tmp_path / "nested" / "resume"

    result = generate_resume_documents(
        improvement,
        profile,
        output_dir,
    )

    assert result.docx_path.exists()
    assert result.pdf_path.exists()


def test_invalid_existing_docx_raises_error(tmp_path):
    improvement, profile = _build_improvement(
        format_choice=ResumeFormatChoice.KEEP_EXISTING,
    )

    invalid_source = tmp_path / "missing.docx"

    try:
        generate_resume_documents(
            improvement,
            profile,
            tmp_path / "output",
            existing_docx_path=invalid_source,
        )
    except DocumentGenerationError:
        pass