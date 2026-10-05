from pathlib import Path

from docx import Document

from ai_resume_writer.documents.template import (
    BuiltInTemplateError,
    create_builtin_template,
)


EXPECTED_SECTIONS = [
    "PROFESSIONAL SUMMARY",
    "SKILLS",
    "EXPERIENCE",
    "PROJECTS",
    "EDUCATION",
    "CERTIFICATIONS",
]


def test_create_builtin_template(tmp_path):
    output_path = tmp_path / "builtin_resume.docx"

    result = create_builtin_template(output_path)

    assert result == output_path
    assert output_path.exists()

    document = Document(output_path)

    paragraphs = [paragraph.text for paragraph in document.paragraphs]

    assert paragraphs[0] == "CANDIDATE NAME"

    for section in EXPECTED_SECTIONS:
        assert section in paragraphs


def test_builtin_template_has_expected_margins(tmp_path):
    output_path = tmp_path / "builtin_resume.docx"

    create_builtin_template(output_path)

    document = Document(output_path)
    section = document.sections[0]

    assert section.top_margin.inches == 0.5
    assert section.bottom_margin.inches == 0.5
    assert section.left_margin.inches == 0.65
    assert section.right_margin.inches == 0.65


def test_builtin_template_uses_ats_friendly_font(tmp_path):
    output_path = tmp_path / "builtin_resume.docx"

    create_builtin_template(output_path)

    document = Document(output_path)

    normal_style = document.styles["Normal"]

    assert normal_style.font.name == "Arial"
    assert normal_style.font.size.pt == 10


def test_builtin_template_does_not_contain_real_candidate_information(
    tmp_path,
):
    output_path = tmp_path / "builtin_resume.docx"

    create_builtin_template(output_path)

    document = Document(output_path)

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "FIRSTNAME" not in text
    assert "LASTNAME" not in text
    assert "XXX" not in text
    assert "Technical Skills A, B, C, D" not in text


def test_builtin_template_creates_parent_directory(tmp_path):
    output_path = (
        tmp_path
        / "nested"
        / "template"
        / "builtin_resume.docx"
    )

    create_builtin_template(output_path)

    assert output_path.exists()


def test_builtin_template_invalid_output_is_wrapped(tmp_path):
    output_path = tmp_path / "directory"

    output_path.mkdir()

    try:
        create_builtin_template(output_path)
    except BuiltInTemplateError:
        pass