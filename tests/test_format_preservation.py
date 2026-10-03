from pathlib import Path

import pytest
from docx import Document

from ai_resume_writer.documents.preservation import (
    DocumentPreservationError,
    preserve_docx_format,
)


FIXTURE_PATH = Path("tests/fixtures/sample_resume.docx")


def test_preserve_docx_format(tmp_path):
    output_path = tmp_path / "preserved_resume.docx"

    result = preserve_docx_format(
        FIXTURE_PATH,
        output_path,
    )

    assert result == output_path
    assert output_path.exists()

    document = Document(output_path)

    assert len(document.paragraphs) == 7
    assert len(document.tables) == 0

    assert [paragraph.text for paragraph in document.paragraphs] == [
        "Test Candidate",
        "PROFESSIONAL SUMMARY",
        "Data Scientist with Python experience.",
        "SKILLS",
        "Python, SQL, Pandas",
        "EXPERIENCE",
        "Data Scientist - Test Company",
    ]

    assert all(
        paragraph.style.name == "Normal"
        for paragraph in document.paragraphs
    )


def test_preserve_docx_format_creates_parent_directory(tmp_path):
    output_path = (
        tmp_path
        / "nested"
        / "output"
        / "preserved_resume.docx"
    )

    preserve_docx_format(
        FIXTURE_PATH,
        output_path,
    )

    assert output_path.exists()


def test_missing_source_raises_error(tmp_path):
    source_path = tmp_path / "missing.docx"
    output_path = tmp_path / "output.docx"

    with pytest.raises(DocumentPreservationError):
        preserve_docx_format(
            source_path,
            output_path,
        )


def test_non_docx_source_raises_error(tmp_path):
    source_path = tmp_path / "resume.pdf"
    source_path.write_bytes(b"not a real pdf")

    output_path = tmp_path / "output.docx"

    with pytest.raises(DocumentPreservationError):
        preserve_docx_format(
            source_path,
            output_path,
        )