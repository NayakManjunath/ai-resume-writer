from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

from ai_resume_writer.contracts.models import (
    CandidateProfile,
    ResumeFormatChoice,
    ResumeImprovement,
)


class DocumentGenerationError(ValueError):
    """Raised when a resume document cannot be generated."""


@dataclass(frozen=True)
class PageValidationResult:
    """Result of deterministic PDF page-count validation."""

    page_count: int
    expected_max_pages: int | None
    is_valid: bool


@dataclass(frozen=True)
class DocumentGenerationResult:
    """Generated resume document paths and page validation result."""

    docx_path: Path
    pdf_path: Path
    page_validation: PageValidationResult


def generate_resume_documents(
    improvement: ResumeImprovement,
    candidate_profile: CandidateProfile,
    output_dir: Path,
    existing_docx_path: Path | None = None,
) -> DocumentGenerationResult:
    """Generate DOCX/PDF resume documents and validate PDF page count."""
    output = Path(output_dir)

    try:
        output.mkdir(parents=True, exist_ok=True)

        docx_path = output / "resume.docx"
        pdf_path = output / "resume.pdf"

        if (
            improvement.format_choice
            == ResumeFormatChoice.KEEP_EXISTING
            and existing_docx_path is not None
        ):
            document = Document(existing_docx_path)
            _clear_document(document)
        else:
            document = Document()

        _configure_document(document)

        _write_docx_content(
            document,
            improvement,
        )

        document.save(docx_path)

        _write_pdf_content(
            pdf_path,
            improvement,
        )

        page_count = _count_pdf_pages(pdf_path)
        expected_max_pages = _expected_max_pages(candidate_profile)

        validation = PageValidationResult(
            page_count=page_count,
            expected_max_pages=expected_max_pages,
            is_valid=(
                expected_max_pages is None
                or page_count <= expected_max_pages
            ),
        )

        return DocumentGenerationResult(
            docx_path=docx_path,
            pdf_path=pdf_path,
            page_validation=validation,
        )

    except DocumentGenerationError:
        raise
    except Exception as exc:
        raise DocumentGenerationError(
            "Unable to generate resume documents."
        ) from exc


def _configure_document(document: Document) -> None:
    """Apply the deterministic resume document configuration."""
    section = document.sections[0]

    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10)


def _write_docx_content(
    document: Document,
    improvement: ResumeImprovement,
) -> None:
    """Write structured resume sections into the DOCX document."""
    _add_name(document, improvement.original_profile.name)

    for section_name, values in improvement.improved_sections.items():
        if not values:
            continue

        _add_section_heading(document, section_name)

        for value in values:
            paragraph = document.add_paragraph(style="Normal")
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.add_run(value)


def _add_name(
    document: Document,
    name: str | None,
) -> None:
    """Add the candidate name when available."""
    if not name or not name.strip():
        return

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(name.strip())
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(16)


def _add_section_heading(
    document: Document,
    title: str,
) -> None:
    """Add a consistently formatted section heading."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)

    run = paragraph.add_run(title.upper())
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10)


def _clear_document(document: Document) -> None:
    """Remove existing body content while retaining document sections."""
    body = document._element.body

    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue

        body.remove(child)


def _write_pdf_content(
    output_path: Path,
    improvement: ResumeImprovement,
) -> None:
    """Generate a PDF from the structured resume content."""
    document = SimpleDocTemplate(
        str(output_path),
        pagesize=LETTER,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )

    styles = getSampleStyleSheet()

    name_style = styles["Title"]
    name_style.fontName = "Helvetica-Bold"
    name_style.fontSize = 16
    name_style.leading = 18
    name_style.alignment = 1

    heading_style = styles["Heading2"]
    heading_style.fontName = "Helvetica-Bold"
    heading_style.fontSize = 10
    heading_style.leading = 12
    heading_style.spaceBefore = 6
    heading_style.spaceAfter = 2

    body_style = styles["BodyText"]
    body_style.fontName = "Helvetica"
    body_style.fontSize = 9
    body_style.leading = 11
    body_style.spaceAfter = 2

    story = []

    name = improvement.original_profile.name

    if name and name.strip():
        story.append(Paragraph(name.strip(), name_style))
        story.append(Spacer(1, 4))

    for section_name, values in improvement.improved_sections.items():
        if not values:
            continue

        story.append(
            Paragraph(
                section_name.upper(),
                heading_style,
            )
        )

        items = [
            ListItem(
                Paragraph(value, body_style),
                leftIndent=10,
            )
            for value in values
            if value.strip()
        ]

        if items:
            story.append(
                ListFlowable(
                    items,
                    bulletType="bullet",
                    leftIndent=14,
                )
            )

    document.build(story)


def _count_pdf_pages(pdf_path: Path) -> int:
    """Count rendered pages in the generated PDF."""
    from pypdf import PdfReader

    reader = PdfReader(pdf_path)
    return len(reader.pages)


def _expected_max_pages(
    candidate_profile: CandidateProfile,
) -> int | None:
    """Return the default maximum page count for the candidate."""
    experience_years = candidate_profile.experience_years

    if experience_years is not None and experience_years <= 4:
        return 1

    return None