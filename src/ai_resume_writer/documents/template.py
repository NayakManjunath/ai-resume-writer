from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


class BuiltInTemplateError(ValueError):
    """Raised when the built-in resume template cannot be created."""


def create_builtin_template(output_path: Path) -> Path:
    """Create the professional ATS-friendly built-in resume template."""
    output = Path(output_path)

    try:
        document = Document()

        section = document.sections[0]
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

        normal_style = document.styles["Normal"]
        normal_style.font.name = "Arial"
        normal_style.font.size = Pt(10)

        _add_header_placeholder(document)
        _add_section_heading(document, "PROFESSIONAL SUMMARY")
        _add_section_heading(document, "SKILLS")
        _add_section_heading(document, "EXPERIENCE")
        _add_section_heading(document, "PROJECTS")
        _add_section_heading(document, "EDUCATION")
        _add_section_heading(document, "CERTIFICATIONS")

        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(output)

    except Exception as exc:
        raise BuiltInTemplateError(
            f"Unable to create built-in resume template: {output}"
        ) from exc

    return output


def _add_header_placeholder(document: Document) -> None:
    """Add a neutral header structure without candidate information."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run("CANDIDATE NAME")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(16)


def _add_section_heading(document: Document, title: str) -> None:
    """Add a consistently formatted resume section heading."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)

    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10)