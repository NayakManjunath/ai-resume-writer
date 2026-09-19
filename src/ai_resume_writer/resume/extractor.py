from pathlib import Path

from docx import Document
from pypdf import PdfReader

from ai_resume_writer.contracts.models import ResumeDocument


class ResumeExtractionError(ValueError):
    """Raised when resume text extraction fails."""


SECTION_NAMES = {
    "summary",
    "professional summary",
    "objective",
    "experience",
    "work experience",
    "professional experience",
    "education",
    "skills",
    "technical skills",
    "projects",
    "certifications",
    "achievements",
    "leadership",
    "extra-curricular activities",
}


def _clean_text(text: str) -> str:
    """Normalize extracted document text."""

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    return "\n".join(lines)


def _extract_pdf_text(file_bytes: bytes) -> str:
    """Extract text from PDF bytes."""

    try:
        from io import BytesIO

        reader = PdfReader(BytesIO(file_bytes))

        pages = []

        for page in reader.pages:
            page_text = page.extract_text() or ""
            pages.append(page_text)

        return _clean_text("\n".join(pages))

    except Exception as exc:
        raise ResumeExtractionError(
            "Unable to extract text from the PDF resume."
        ) from exc


def _extract_docx_text(file_bytes: bytes) -> str:
    """Extract text from DOCX bytes."""

    try:
        from io import BytesIO

        document = Document(BytesIO(file_bytes))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        return _clean_text("\n".join(paragraphs))

    except Exception as exc:
        raise ResumeExtractionError(
            "Unable to extract text from the DOCX resume."
        ) from exc


def _detect_sections(text: str) -> list[str]:
    """Detect common resume section headings."""

    detected_sections = []

    for line in text.splitlines():
        normalized = line.strip().lower()

        if normalized in SECTION_NAMES:
            detected_sections.append(line.strip())

    return detected_sections


def extract_resume(
    file_name: str,
    file_bytes: bytes,
) -> ResumeDocument:
    """
    Extract text and basic section headings from a resume.

    Supports PDF and DOCX files.
    """

    if not file_name or not file_name.strip():
        raise ResumeExtractionError(
            "Resume file name cannot be empty."
        )

    if not file_bytes:
        raise ResumeExtractionError(
            "Resume file cannot be empty."
        )

    extension = Path(file_name).suffix.lower()

    if extension == ".pdf":
        text = _extract_pdf_text(file_bytes)

    elif extension == ".docx":
        text = _extract_docx_text(file_bytes)

    else:
        raise ResumeExtractionError(
            "Unsupported resume format. "
            "Supported formats: .pdf, .docx"
        )

    if not text:
        raise ResumeExtractionError(
            "No usable text could be extracted from the resume."
        )

    sections = _detect_sections(text)

    return ResumeDocument(
        file_name=file_name,
        file_type=extension.lstrip("."),
        raw_text=text,
        sections=sections,
    )